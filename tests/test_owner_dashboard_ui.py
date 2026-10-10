from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import pandas as pd
from openpyxl import load_workbook
from owner_dashboard_ui import activity, p95, in_period, workbook_bytes
from site_analytics import record_session, sessions, record_browser_time


class OwnerAnalyticsTests(unittest.TestCase):
    def test_sessions_are_idempotent_in_isolated_storage(self):
        with TemporaryDirectory() as folder:
            db = Path(folder) / "qa.sqlite3"
            record_session("anonymous-session", db)
            record_browser_time("anonymous-session", 2.5, db)
            record_browser_time("anonymous-session", 99, db)
            record_session("anonymous-session", db)
            rows = sessions(db)
            self.assertEqual(len(rows), 1)
            self.assertEqual(set(rows[0]), {"id", "started_at", "browser_seconds"})
            self.assertEqual(float(rows[0]["browser_seconds"]), 2.5)

    def test_chart_series_separate_website_sessions_and_mira_replies(self):
        rows = [dict(**{"Inquiry ID": str(i), "Conversation ID": "one", "Timestamp (Asia/Kolkata)": "2026-10-10", "Assistant response": "Reply"}) for i in range(2)]
        site, mira = activity(rows + [rows[0]], [{"id": "a", "started_at": "2026-10-10"}]*2)
        self.assertEqual(site, [{"Date": "2026-10-10", "Sessions": 1}])
        self.assertEqual({r["Series"]: r["Count"] for r in mira}, {"Conversations": 1, "Replies": 2})

    def test_missing_timing_is_not_a_fabricated_zero(self):
        self.assertEqual(p95([None, "", "invalid"]), "Not measured")
        self.assertEqual(p95([1, 2, 3, 4]), "4.00 s")
        self.assertTrue(in_period("2026-10-09T20:00:00Z", date(2026,10,10), date(2026,10,10)))

    def test_export_has_four_views_and_escapes_formulas(self):
        sheets = {name: pd.DataFrame({"text": ["=HYPERLINK(\"bad\")"]}) for name in ("Inquiries", "Mira work", "Feedback", "Learning library")}
        book = load_workbook(BytesIO(workbook_bytes(sheets)))
        self.assertEqual(book.sheetnames, list(sheets))
        self.assertNotEqual(book["Feedback"]["A2"].data_type, "f")
        book.close()

    def test_ui_has_four_views_and_preserves_learning_controls(self):
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_file(str(Path(__file__).with_name("owner_dashboard_sandbox.py")), default_timeout=60)
        app.session_state.qa_inquiry_rows = [{"Inquiry ID":"qa", "Conversation ID":"qa", "Timestamp (Asia/Kolkata)":"2026-10-10", "Assistant response":"Test reply", "Mira response seconds":2}]
        app.session_state.qa_sessions = [{"id":"test", "started_at":"2026-10-10"}]
        app.session_state.qa_deliveries = [
            dict(conversation_id="qa", channel="email", status="scheduled", sent_count=2, failed_recent=0, updated_at="2026-10-10"),
            dict(conversation_id="qa", channel="email", status="scheduled", sent_count=3, failed_recent=1, updated_at="2026-10-10")]
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(next(m.value for m in app.metric if m.label == "Follow-up acceptance"), "50%")
        self.assertTrue(any("Emails needing review: 1" in m.value for m in app.markdown))
        self.assertEqual([t.label for t in app.tabs], ["Inquiries", "Mira responses & actions", "Feedback", "Learning library"])
        self.assertTrue(any(b.label == "Save Learning Library" for b in app.button))
        self.assertTrue(any(b.label == "Save outcome updates" for b in app.button))
        self.assertTrue(any(d.label == "Download four connected record views (.xlsx)" for d in app.get("download_button")))
        selector = next(m for m in app.multiselect if m.label == "KPIs to display")
        self.assertEqual(len(selector.options), 14)
        selector.set_value(["Mira conversations", "Mira replies"]).run()
        self.assertFalse(app.exception)
        self.assertEqual(next(m.value for m in app.metric if m.label == "Mira conversations"), "1")
        self.assertFalse(any(m.label == "Website sessions" for m in app.metric))
        self.assertEqual([t.label for t in app.tabs], ["Inquiries", "Mira responses & actions", "Feedback", "Learning library"])
        next(m for m in app.multiselect if m.label == "KPIs to display").set_value([]).run()
        self.assertFalse(app.exception)
        self.assertTrue(any("Choose at least one KPI" in item.value for item in app.info))

    def test_feedback_draft_is_visible_on_first_render_without_approval(self):
        from streamlit.testing.v1 import AppTest
        from mira_learning_library import active_learning_guidance
        app = AppTest.from_file(str(Path(__file__).with_name("owner_dashboard_sandbox.py")), default_timeout=60)
        app.session_state.qa_inquiry_rows = [{"Inquiry ID":"fb", "Feedback ID":"fb", "Conversation ID":"qa", "Timestamp (Asia/Kolkata)":"2026-10-10", "Feedback category":"Misunderstood my request", "Feedback comment":"Keep my chosen locations."}]
        app.run()
        self.assertFalse(app.exception)
        drafts = app.session_state.qa_rules
        self.assertTrue(any(r["Rule ID"] == "feedback-fb" for r in drafts))
        self.assertTrue(all(r["Status"] == "Draft" for r in drafts))
        self.assertEqual(active_learning_guidance(drafts), [])
        self.assertTrue(any("Rule ID" in frame.value.columns and "feedback-fb" in str(frame.value) for frame in app.dataframe))
