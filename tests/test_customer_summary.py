import ast
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock

from agent_runtime import INDIA_TZ
from io import BytesIO

from openpyxl import load_workbook

from inquiry_log import append_inquiry, build_conversation_transcript_export, customer_summary_rows, read_inquiries, update_conversation_fields


class CustomerSummaryTests(unittest.TestCase):
    def test_latest_interest_and_consent_without_raw_answers(self):
        rows = [
            {"Conversation ID": "one", "Timestamp (Asia/Kolkata)": "2026-10-01", "Customer interest signal": "Interested signal", "Customer budget / price stated": "I want a plot under 30 lakh", "User inquiry": "private answer"},
            {"Conversation ID": "one", "Timestamp (Asia/Kolkata)": "2026-10-02", "Customer interest signal": "Not interested signal", "Follow-up method": "Email", "Follow-up consent timestamp (Asia/Kolkata)": "2026-10-02"},
        ]
        summary, = customer_summary_rows(rows)
        self.assertEqual(summary["Interested in property"], "No")
        self.assertEqual(summary["Not interested in property"], "Yes")
        self.assertEqual(summary["Email follow-ups"], "Yes")
        self.assertEqual(summary["Follow-up methods"], "Email")
        self.assertEqual(summary["Budget"], "30 lakh")
        self.assertNotIn("private answer", str(summary))
        unstated, = customer_summary_rows([{"Conversation ID": "two"}])
        self.assertEqual(unstated["Interest stated"], "No")
        self.assertEqual(unstated["Not interested in property"], "No")

    def test_dashboard_summary_keeps_every_consented_followup_method(self):
        summary, = customer_summary_rows([
            {"Conversation ID": "both", "Timestamp (Asia/Kolkata)": "2026-10-01", "Follow-up method": "Email", "Follow-up consent timestamp (Asia/Kolkata)": "2026-10-01"},
            {"Conversation ID": "both", "Timestamp (Asia/Kolkata)": "2026-10-02", "Follow-up method": "WhatsApp", "Follow-up consent timestamp (Asia/Kolkata)": "2026-10-02"},
        ])
        self.assertEqual(summary["Email follow-ups"], "Yes")
        self.assertEqual(summary["WhatsApp follow-ups"], "Yes")
        self.assertEqual(summary["Follow-up methods"], "Email, WhatsApp")

    def test_end_button_saves_then_clears_chat(self):
        class State(dict):
            __getattr__ = dict.__getitem__
            __setattr__ = dict.__setitem__
        state = State(chat=[{"role": "user", "content": "hello"}], inquiry_conversation_id="one", user_id="test-user", agent_history=["history"])
        tree = ast.parse((Path(__file__).parents[1] / "app.py").read_text(encoding="utf-8-sig"))
        function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "close_conversation_from_button")
        save = Mock()
        sender = Mock(return_value=(0, 0))
        namespace = {"st": SimpleNamespace(session_state=state), "datetime": datetime, "INDIA_TZ": INDIA_TZ,
                     "language": "English", "update_conversation_fields": save, "conversation_transcript": lambda chat: str(chat),
                     "clear_active_preferences": Mock(), "send_chat_end_recommendations": sender,
                     "FOLLOWUP_EMAIL_CONFIG": {}, "current_followup_recommendations": lambda: "Saved shortlist"}
        exec(compile(ast.Module(body=[function], type_ignores=[]), "app.py", "exec"), namespace)
        namespace["close_conversation_from_button"]()
        sender.assert_called_once_with(config={},user_id="test-user",conversation_id="one",recommendations="Saved shortlist")
        self.assertTrue(state.conversation_closed)
        self.assertEqual(state.chat, [])
        self.assertEqual(state.agent_history, [])
        self.assertIn("hello", save.call_args.args[1]["Final conversation transcript"])

    def test_rating_is_kept_in_the_structured_summary(self):
        summary, = customer_summary_rows([
            {"Conversation ID": "rated", "Mira performance rating (1-5)": 4, "User satisfaction rating (1-5)": 4},
        ])
        self.assertEqual(summary["Mira performance rating (1-5)"], "4")
        self.assertEqual(summary["User satisfaction rating (1-5)"], "4")

    def test_complete_chat_export_includes_private_transcript_and_rating(self):
        exported = build_conversation_transcript_export([
            {
                "Inquiry ID": "NMI-RATED-1", "Conversation ID": "rated-chat", "Customer name": "Sam",
                "Timestamp (Asia/Kolkata)": "2026-10-08T12:00:00+05:30", "User inquiry": "Need a 2BHK",
                "Assistant response": "I can help with that.", "Conversation status": "Ended",
                "Mira performance rating (1-5)": 5, "User satisfaction rating (1-5)": 5,
            },
        ], "October 2026")
        workbook = load_workbook(BytesIO(exported), data_only=True)
        sheet = workbook["Complete chats"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index("Mira performance rating (1-5)") + 1).value, "5")
        transcript = sheet.cell(2, headers.index("Complete chat transcript") + 1).value
        self.assertIn("Customer: Need a 2BHK", transcript)
        self.assertIn("Mira: I can help with that.", transcript)
        self.assertIn("Export details", workbook.sheetnames)
        workbook.close()

    def test_rating_updates_every_saved_turn_in_the_private_workbook(self):
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "inquiries.xlsx"
            append_inquiry({"Inquiry ID": "one", "Conversation ID": "shared", "User inquiry": "First"}, path)
            append_inquiry({"Inquiry ID": "two", "Conversation ID": "shared", "User inquiry": "Second"}, path)
            updated = update_conversation_fields("shared", {
                "Mira performance rating (1-5)": 3,
                "User satisfaction rating (1-5)": 3,
            }, path)
            self.assertEqual(updated, 2)
            rows = read_inquiries(path)
            self.assertEqual({row["Mira performance rating (1-5)"] for row in rows}, {"3"})
            self.assertEqual({row["User satisfaction rating (1-5)"] for row in rows}, {"3"})
