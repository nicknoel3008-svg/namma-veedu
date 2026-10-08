import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pandas as pd
import csv
import json
import sys
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from agent import parse_request
from agent_policy import AGENT_SYSTEM_POLICY, FRIENDLY_CMDA_NOTE, requires_user_approval
from agent_runtime import dispatch_tool, run_openai_agent
from area_utils import convert_area, price_per_sqm
from property_search import load_properties, search_properties
from mira_dialogue_libraries import local_dialogue_reply, local_land_conversation_reply
from inquiry_log import build_inquiry_export
from openpyxl import load_workbook


class PropertyFinderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load_properties()

    def test_square_foot_conversion(self):
        self.assertAlmostEqual(convert_area(1000, "sq ft")["sq_m"], 92.903, places=2)
        self.assertAlmostEqual(convert_area(1, "hectares")["sq_ft"], 107_639.104, places=2)

    def test_csv_calculation_columns_load_as_numbers(self):
        self.assertTrue(pd.api.types.is_numeric_dtype(self.data["area_sqm"]))
        self.assertTrue(pd.api.types.is_numeric_dtype(self.data["price_per_sqm_inr"]))
        self.assertTrue((self.data["price_inr"].dropna() > 0).all())
        self.assertFalse(pd.to_numeric(self.data["price_display"], errors="coerce").eq(0).any())
        self.assertTrue((self.data["area_sqm"].dropna() > 0).all())

    def test_price_per_square_metre(self):
        self.assertAlmostEqual(price_per_sqm(1_000_000, 100, "sq m"), 10_000)

    def test_parse_city_budget_type_and_auction(self):
        intent = parse_request("Find flat auctions in Chennai under ₹50 lakh", self.data)
        self.assertEqual(intent.location, "Chennai")
        self.assertEqual(intent.property_type, "Flat")
        self.assertEqual(intent.status, "Auction")
        self.assertEqual(intent.max_budget, 5_000_000)

    def test_filter_location_type_and_budget(self):
        found = search_properties(self.data, location="Chennai", property_type="Flat", max_budget=10_000_000)
        self.assertTrue(all(found["property_type"] == "Flat"))
        self.assertTrue(all(found["price_inr"].dropna() <= 10_000_000))
        self.assertTrue(found["city"].str.contains("Chennai", case=False, na=False).any())

    def test_location_search_does_not_confuse_porur_with_thiruporur(self):
        sample = pd.DataFrame([
            {"city": "Chennai", "district": "Chennai", "locality": "PORUR", "location": "", "property_type": "Flat", "listing_status": "Existing sale", "price_inr": 5_000_000, "area_sqm": 80, "price_per_sqm_inr": 62_500, "source_status": "verified", "is_auction_open": False},
            {"city": "Chennai", "district": "Chengalpattu", "locality": "Thiruporur", "location": "", "property_type": "Flat", "listing_status": "Existing sale", "price_inr": 5_000_000, "area_sqm": 80, "price_per_sqm_inr": 62_500, "source_status": "verified", "is_auction_open": False},
        ])
        found = search_properties(sample, location="Porur")
        self.assertEqual(found["locality"].tolist(), ["PORUR"])

    def test_multi_select_filters_combine_categories_and_include_selected_values(self):
        sample = pd.DataFrame([
            {"city":"Chennai", "district":"Chennai", "locality":"A", "location":"", "property_type":"Plot", "listing_status":"Existing sale", "price_inr":1_000_000, "area_sqm":100, "price_per_sqm_inr":10_000, "source_status":"verified", "is_auction_open":False},
            {"city":"Madurai", "district":"Madurai", "locality":"B", "location":"", "property_type":"Flat", "listing_status":"Project reference", "price_inr":2_000_000, "area_sqm":80, "price_per_sqm_inr":25_000, "source_status":"verified", "is_auction_open":False},
            {"city":"Coimbatore", "district":"Coimbatore", "locality":"C", "location":"", "property_type":"House", "listing_status":"Existing sale", "price_inr":3_000_000, "area_sqm":120, "price_per_sqm_inr":25_000, "source_status":"verified", "is_auction_open":False},
            {"city":"Chennai", "district":"Chennai", "locality":"D", "location":"", "property_type":"Plot", "listing_status":"Auction ended", "price_inr":500_000, "area_sqm":90, "price_per_sqm_inr":5_555, "source_status":"verified", "is_auction_open":False},
        ])
        found = search_properties(
            sample,
            location=["Chennai", "Madurai"],
            property_type=["Plot", "Flat"],
            status=["Existing sale", "Project reference"],
        )
        self.assertEqual(set(found["locality"]), {"A", "B"})

    def test_missing_area_rows_are_kept_without_size_filter(self):
        data = pd.DataFrame([{
            "city":"Chennai", "district":"Chennai", "locality":"", "location":"",
            "property_type":"Plot", "listing_status":"Existing sale", "price_inr":1_000_000,
            "area_sqm":float("nan"), "price_per_sqm_inr":float("nan"), "source_status":"snapshot",
            "auction_end":"", "is_auction_open":False,
        }])
        self.assertEqual(len(search_properties(data)), 1)
        self.assertEqual(len(search_properties(data, min_area_sqm=1)), 0)

    def test_listing_source_links_and_sanitized_fields(self):
        source_file = ROOT / "data" / "properties.csv"
        with source_file.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            rows = list(reader)
            self.assertTrue(all(not row["source_url"] or row["source_url"].startswith("https://") for row in rows))
            self.assertNotIn("borrower_name", reader.fieldnames)
            self.assertNotIn("borrower_address", reader.fieldnames)

    def test_tamil_search_and_approval_register_is_not_a_sale_listing(self):
        intent = parse_request("சென்னையில் ₹30 லட்சத்திற்குள் மனை காட்டு", self.data)
        self.assertEqual(intent.location, "Chennai")
        self.assertEqual(intent.property_type, "Plot")
        self.assertEqual(intent.max_budget, 3_000_000)
        regular = search_properties(self.data)
        self.assertFalse(regular["listing_status"].eq("Approval record").any())

    def test_auction_and_loan_intents_stay_separate(self):
        auction = parse_request("Show bank auction houses in Chennai", self.data)
        self.assertEqual(auction.status, "Auction")
        self.assertFalse(auction.help_with_loan)
        loan = parse_request("Estimate my home loan EMI", self.data)
        self.assertTrue(loan.help_with_loan)

    def test_land_safety_library_does_not_override_an_explicit_plot_search(self):
        reply = local_land_conversation_reply(
            "I need a plot in Chennai under 40 lakh. Show me a few options.", [], "English"
        )
        self.assertIsNone(reply)

    def test_disappointed_customer_reply_respects_the_requested_property_type(self):
        history = [{"role": "assistant", "mode": "results", "records": [{"title": "Saved plot"}]}]
        reply = local_dialogue_reply(
            "I am disappointed. You are not listening; I asked for a plot, not a flat.",
            history,
            "English",
        )
        self.assertIn("plots, not flats", reply)
        self.assertIn("sorry", reply.casefold())

    def test_owner_export_contains_selected_rows_and_period_metadata(self):
        exported = build_inquiry_export(
            [{"Inquiry ID": "NMI-TEST-1", "Customer name": "Sample", "User inquiry": "Find a plot"}],
            "October 2026",
            "2026-10-01",
            "2026-10-31",
        )
        workbook = load_workbook(BytesIO(exported), data_only=True)
        inquiries = workbook["Inquiries"]
        details = workbook["Export details"]
        headers = [cell.value for cell in inquiries[1]]
        self.assertEqual(inquiries.cell(2, headers.index("Inquiry ID") + 1).value, "NMI-TEST-1")
        self.assertEqual(inquiries.cell(2, headers.index("Customer name") + 1).value, "Sample")
        self.assertNotIn("User inquiry", headers)
        self.assertNotIn("Assistant response", headers)
        self.assertNotIn("Final conversation transcript", headers)
        self.assertEqual(inquiries.cell(2, headers.index("Email follow-ups") + 1).value, "No")
        self.assertEqual(details.cell(1, 2).value, "October 2026")
        self.assertEqual(details.cell(4, 2).value, 1)
        self.assertIn("Mira responses", workbook.sheetnames)
        responses = workbook["Mira responses"]
        response_headers = [cell.value for cell in responses[1]]
        self.assertIn("Mira response", response_headers)
        self.assertEqual(responses.max_row, 1)
        workbook.close()

    def test_owner_export_keeps_mira_responses_out_of_summary_sheet(self):
        exported = build_inquiry_export(
            [{"Inquiry ID": "NMI-TEST-2", "Conversation ID": "C-2", "Timestamp (Asia/Kolkata)": "2026-10-08T12:00:00+05:30", "Language": "English", "Response type": "results", "User inquiry": "Show me flats", "Assistant response": "I found two saved flats."}],
            "October 2026", "2026-10-01", "2026-10-31",
        )
        workbook = load_workbook(BytesIO(exported), data_only=True)
        summary_headers = [cell.value for cell in workbook["Inquiries"][1]]
        response_headers = [cell.value for cell in workbook["Mira responses"][1]]
        self.assertNotIn("Assistant response", summary_headers)
        self.assertEqual(workbook["Mira responses"].cell(2, response_headers.index("Mira response") + 1).value, "I found two saved flats.")
        self.assertNotIn("Show me flats", [cell.value for cell in workbook["Mira responses"][2]])
        workbook.close()

    def test_sources_contain_auction_and_loan_pages(self):
        source_file = ROOT / "data" / "sources.csv"
        with source_file.open(encoding="utf-8", newline="") as handle:
            sources = list(csv.DictReader(handle))
        self.assertTrue(any(row["category"] == "Auction" for row in sources))
        self.assertTrue(any(row["category"] == "Loan" for row in sources))
        self.assertTrue(all(row["url"].startswith("https://") for row in sources))

    def test_agent_trust_policy_and_confirmation_boundaries(self):
        self.assertIn("CMDA planning-permission records do not establish property value", AGENT_SYSTEM_POLICY)
        self.assertIn("current price", FRIENDLY_CMDA_NOTE)
        self.assertTrue(requires_user_approval("send_email"))
        self.assertTrue(requires_user_approval("create_external_reminder"))
        self.assertFalse(requires_user_approval("search_saved_properties"))

    def test_agent_tool_searches_keep_property_auctions_and_cmda_separate(self):
        sources = pd.read_csv(ROOT / "data" / "sources.csv", keep_default_na=False)
        normal = dispatch_tool("search_saved_properties", {
            "city": "Chennai", "property_type": "Any", "bedrooms": 0, "record_kind": "Any",
            "max_budget_inr": 0, "min_area_sqm": 0, "limit": 8,
        }, self.data, sources)
        self.assertEqual(normal.display_kind, "properties")
        self.assertFalse(any(row.get("listing_status") in ("Auction", "Auction ended", "Approval record") for row in normal.data["records"]))

        auctions = dispatch_tool("search_bank_auctions", {
            "city": "Chennai", "property_type": "Any", "include_ended": True, "limit": 8,
        }, self.data, sources)
        self.assertEqual(auctions.display_kind, "auctions")
        self.assertTrue(all(row.get("listing_status") in ("Auction", "Auction ended") for row in auctions.data["records"]))

        cmda = dispatch_tool("search_cmda_approvals", {
            "city": "Chennai", "approval_number": "", "limit": 5,
        }, self.data, sources)
        self.assertEqual(cmda.display_kind, "cmda")
        self.assertTrue(all("price_inr" not in row and "availability" not in row for row in cmda.data["records"]))

    def test_undated_unavailable_auctions_require_explicit_selection(self):
        sample = self.data[self.data["listing_status"].eq("Auction date not listed")].head(1).copy()
        city = str(sample.iloc[0]["city"])
        regular_auction_search = search_properties(sample, location=city, status="Auction")
        self.assertTrue(regular_auction_search.empty)

        undated = search_properties(sample, location=city, status=["Auction date not listed"])
        self.assertGreater(len(undated), 0)
        self.assertTrue(undated["auction_end"].eq("").all())

        sources = pd.read_csv(ROOT / "data" / "sources.csv", keep_default_na=False)
        result = dispatch_tool("search_bank_auctions", {
            "city": city, "property_type": "Any", "include_ended": False, "limit": 3,
        }, sample, sources)
        self.assertEqual(result.data["count"], 0)
        self.assertGreater(result.data["undated_count"], 0)
        self.assertEqual(result.data["records"], [])

    def test_followup_tool_only_proposes_and_missing_key_does_not_call_api(self):
        sources = pd.read_csv(ROOT / "data" / "sources.csv", keep_default_na=False)
        proposal = dispatch_tool("propose_in_app_followup", {
            "due_at": "2099-01-01T10:00:00+05:30", "preferences_summary": "Plots near Chennai",
        }, self.data, sources)
        self.assertEqual(proposal.data["status"], "waiting_for_user_confirmation")
        empty = run_openai_agent(text="hello", history=[], properties=self.data, sources=sources, api_key="")
        self.assertEqual(empty.error, "missing_api_key")

    def test_openai_tool_loop_executes_only_registered_tools(self):
        sources = pd.read_csv(ROOT / "data" / "sources.csv", keep_default_na=False)

        class FakeResponse:
            def __init__(self, output, output_text=""):
                self.output = output
                self.output_text = output_text

        class FakeCall:
            type = "function_call"
            name = "search_saved_properties"
            call_id = "call_search_1"
            arguments = json.dumps({
                "city": "Chennai", "property_type": "Any", "bedrooms": 0, "record_kind": "Any",
                "max_budget_inr": 0, "min_area_sqm": 0, "limit": 3,
            })

            def model_dump(self, exclude_none=True):
                return {"type": self.type, "name": self.name, "call_id": self.call_id, "arguments": self.arguments}

        class FakeResponses:
            def __init__(self):
                self.calls = []
                self.responses = [FakeResponse([FakeCall()]), FakeResponse([], "I found saved options in Chennai.")]

            def create(self, **kwargs):
                self.calls.append(kwargs)
                return self.responses.pop(0)

        fake_responses = FakeResponses()
        fake_client = SimpleNamespace(responses=fake_responses)
        fake_openai = SimpleNamespace(OpenAI=lambda api_key, **kwargs: fake_client)
        with patch.dict(sys.modules, {"openai": fake_openai}):
            turn = run_openai_agent(text="Find properties in Chennai", history=[], properties=self.data, sources=sources, api_key="test-only")

        self.assertEqual(turn.text, "I found saved options in Chennai.")
        self.assertEqual(len(turn.tool_results), 1)
        self.assertFalse(turn.error)
        self.assertTrue(any(item.get("type") == "function_call_output" for item in fake_responses.calls[1]["input"]))


if __name__ == "__main__":
    unittest.main()
