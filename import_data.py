"""Copy and normalize the user's TN Properties files into the app's CSV.

Only public listing/project fields are exported. Borrower/owner names and
addresses, bank branch addresses, photos and IDs that are not user-facing are
intentionally omitted.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
import re
from pathlib import Path

import openpyxl

from area_utils import convert_area, price_per_sqm

DEFAULT_SOURCE = Path.home() / "OneDrive" / "Documents" / "TN Properties"
OUT = Path(__file__).parent / "data" / "properties.csv"
CHECKED = date.today().isoformat()
FIELDS = [
    "property_id", "title", "property_type", "listing_status", "city", "district",
    "locality", "price_inr", "price_display", "area_value", "area_unit", "area_display",
    "area_sqm", "price_per_sqm_inr", "bedrooms", "bank_name", "ownership_type",
    "possession_type", "auction_start", "auction_end", "emd_deadline", "inspection_time",
    "approval_number", "rera_number", "availability", "amenities", "price_basis",
    "facing", "floor_number", "car_parking", "furnished_status", "project_name",
    "action_type", "loan_available", "loan_percentage",
    "source_name", "source_url", "source_date", "date_checked", "source_status", "notes",
]


def clean(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def number(value):
    if value is None:
        return None
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", str(value).replace("₹", ""))
    if not match:
        return None
    try:
        return float(match.group(0).replace(",", ""))
    except ValueError:
        return None


def normalize_type(raw: str) -> str | None:
    text = raw.casefold()
    if any(term in text for term in ("plot", "land", "site")):
        return "Plot"
    if any(term in text for term in ("flat", "apartment", "studio")):
        return "Flat"
    if any(term in text for term in ("house", "villa", "bungalow", "farm house")):
        return "House"
    return None


def normalize_unit(raw: str) -> str:
    unit = clean(raw).lower().replace("²", "2")
    if unit in {"sq feet", "sq. feet", "sq ft", "sqft", "square feet", "square foot"}:
        return "sq ft"
    if unit in {"sq m", "sqm", "square metres", "square meters", "square metre", "square meter"}:
        return "sq m"
    if unit in {"cent", "cents"}:
        return "cents"
    if unit in {"acre", "acres"}:
        return "acres"
    if unit in {"hectare", "hectares", "ha"}:
        return "hectares"
    if unit in {"ground", "grounds"}:
        return "grounds"
    return unit


def add(rows: list[dict], **record):
    row = {field: "" for field in FIELDS}
    row.update({key: clean(value) for key, value in record.items() if key in row})
    price = number(record.get("price_inr"))
    if price is not None and price > 0:
        row["price_inr"] = str(round(price, 2))
    else:
        row["price_inr"] = ""
        if number(record.get("price_display")) == 0:
            row["price_display"] = ""
    area_value = number(record.get("area_value"))
    unit = normalize_unit(record.get("area_unit", ""))
    if area_value is not None and area_value > 0 and unit:
        row["area_value"] = str(area_value)
        row["area_unit"] = unit
        try:
            row["area_sqm"] = str(round(convert_area(area_value, unit)["sq_m"], 3))
            if price is not None:
                row["price_per_sqm_inr"] = str(round(price_per_sqm(price, area_value, unit), 2))
        except ValueError:
            row["area_sqm"] = ""
            row["price_per_sqm_inr"] = ""
    elif area_value is None or area_value <= 0:
        row["area_value"] = ""
        row["area_unit"] = ""
    rows.append(row)


def export_auction_csv(folder: Path, rows: list[dict]):
    path = folder / "tamil_nadu_properties.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for source in csv.DictReader(handle):
            prop_type = normalize_type(source.get("property_sub_type") or source.get("asset_type") or "")
            if not prop_type:
                continue
            detail_id = clean(source.get("property_detail_id"))
            event_end = clean(source.get("auction_end"))
            is_flagged_available = clean(source.get("auction_available")).lower() == "true"
            end_date = event_end[:10]
            try:
                is_open = is_flagged_available and date.fromisoformat(end_date) >= date.today()
            except ValueError:
                is_open = False
            listing_status = (
                "Auction" if is_open else "Auction ended" if is_flagged_available
                else "Auction date not listed"
            )
            raw_area = source.get("area")
            unit = normalize_unit(source.get("measurement") or source.get("measurement_abbreviation") or "")
            add(
                rows,
                property_id=f"AUC-{detail_id}",
                title=source.get("property_heading") or source.get("property_title") or source.get("property_sub_type") or "Bank property auction",
                property_type=prop_type,
                listing_status=listing_status,
                city=source.get("city"), district=source.get("district"), locality=source.get("locality"),
                price_inr=source.get("property_price"), price_display=source.get("property_price"),
                area_value=raw_area, area_unit=unit,
            area_display=(
                f"{clean(raw_area)} {unit}".strip()
                if number(raw_area) and number(raw_area) > 0 and unit in {
                    "sq ft", "sq m", "cents", "acres", "hectares", "grounds"
                } else ""
            ),
                bank_name=source.get("bank_name"), ownership_type=source.get("ownership_type"),
                possession_type=source.get("possession_type"), auction_start=source.get("auction_start"),
                auction_end=event_end, emd_deadline=source.get("emd_end"), inspection_time=source.get("inspection_start"),
                facing=source.get("facing"), floor_number=source.get("floor_number"),
                car_parking=source.get("car_parking"), furnished_status=source.get("furnished_status"),
                project_name=source.get("project_name"), action_type=source.get("action_type"),
                loan_available=source.get("loan_available"), loan_percentage=source.get("loan_percentage"),
                source_name="User-provided Tamil Nadu auction snapshot",
                source_url="https://baanknet.com/",
                source_date=clean(source.get("created_on"))[:10], date_checked=CHECKED,
                source_status=(
                    "Snapshot import; auction date not listed and availability marked false"
                    if not is_flagged_available else "Snapshot import; not live-verified"
                ),
                notes=("Auction dates, reserve price, possession and title must be checked in the bank's current sale notice. "
                       + ("The source marks this record unavailable and provides no auction date; it is shown only when the user asks to review undated records."
                          if not is_flagged_available else
                          "Auction end date is past; historical record only." if not is_open else
                          "Upcoming in the snapshot; confirm it remains open.")),
            )


def export_nobroker(folder: Path, rows: list[dict]):
    book = openpyxl.load_workbook(folder / "tamil_nadu_properties(NO broker).xlsx", read_only=True, data_only=True)
    sheet = book["Property_Master"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for values in iterator:
        source = dict(zip(headers, values))
        prop_type = normalize_type(source.get("Property_Type_Clean") or source.get("Title") or "")
        if not prop_type:
            continue
        row_id = clean(source.get("Property_ID"))
        sale_price_lakh = number(source.get("Price_Lakhs"))
        add(
            rows,
            property_id=f"SALE-{row_id[:12]}", title=source.get("Title"), property_type=prop_type,
            listing_status="Existing sale", city=source.get("District"), district=source.get("District"),
            locality=source.get("Locality"), price_inr=sale_price_lakh * 100_000 if sale_price_lakh is not None else None,
            price_display=f"₹{clean(source.get('Price_Lakhs'))} lakh (listing price)",
            area_value=source.get("Area_Sqft_Numeric"), area_unit="sq ft",
            area_display=f"{clean(source.get('Area_Sqft_Numeric'))} sq ft" if (number(source.get("Area_Sqft_Numeric")) or 0) > 0 else "",
            bedrooms=source.get("BHK_Numeric"), source_name=source.get("Source") or "NoBroker listing extract",
            source_url=source.get("URL"), source_date=source.get("Extraction_Date"), date_checked=CHECKED,
            source_status="Saved listing snapshot; ownership history not independently verified",
            notes="Treat as an advertised sale listing. Confirm current availability, seller authority, title and encumbrances independently.",
        )


def export_stepsstone(folder: Path, rows: list[dict]):
    book = openpyxl.load_workbook(folder / "StepsStone_Master_Project_Dataset_Field_Verified.xlsx", read_only=True, data_only=True)
    sheet = book["13_Verified_Project_Data"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for index, values in enumerate(iterator, start=1):
        source = dict(zip(headers, values))
        prop_type = normalize_type(source.get("Category") or "")
        if not prop_type:
            continue
        raw_size = clean(source.get("Size_Range_SqFt"))
        exact_size = number(raw_size) if re.fullmatch(r"[\d,.]+", raw_size) else None
        raw_price = clean(source.get("Price"))
        exact_price = number(raw_price) if re.fullmatch(r"[\d,.]+", raw_price) else None
        name = clean(source.get("Project_Name"))
        add(
            rows,
            property_id=f"PROJ-{index:03d}", title=name,
            property_type=prop_type, listing_status="Project reference", city=source.get("Location"),
            district=source.get("Location"), locality=source.get("Location"), price_inr=exact_price,
            price_display=raw_price, area_value=exact_size, area_unit="sq ft" if exact_size else "",
            area_display=f"{raw_size} sq ft" if raw_size else "",
            source_name="StepsStone verified project workbook", source_url=source.get("Official_Project_Source_URL"),
            source_date=source.get("Checked_Date"), date_checked=CHECKED,
            source_status="Project reference; not a confirmed individual available unit",
            notes="Project-level information only. Confirm current inventory and quoted price with the project source. "
                  + clean(source.get("Verification_Notes")),
        )


def export_vgn(folder: Path, rows: list[dict]):
    path = folder / "VGN_MASTER_DATASET_PRICE_AREA_AVAILABILITY_AUDITED_2026-10-06.xlsx"
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["01_Master_Project_Inventory"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for values in iterator:
        source = dict(zip(headers, values))
        title = clean(source.get("Project_Name"))
        prop_type = normalize_type(clean(source.get("Project_Type") or ""))
        if not title or not prop_type:
            continue
        raw_price = clean(source.get("Price"))
        price_unit = clean(source.get("Price_Unit")).casefold()
        price_value = number(raw_price)
        raw_price_folded = raw_price.casefold()
        is_area_rate = bool(re.search(
            r"(?:/|\bper\s*)(?:sq\.?\s*ft|sqft|square\s*feet|sq\.?\s*m|sqm|square\s*met(?:er|re)s?|acre|acres|cent|cents|ground|grounds)\b",
            raw_price_folded,
        ))
        if price_value is None or "not published" in price_unit or is_area_rate:
            price = None
        elif re.search(r"\b(?:crore|crores|cr)\b", raw_price_folded):
            price = price_value * 10_000_000
        elif re.search(r"\b(?:lakh|lakhs|lac|lacs)\b", raw_price_folded):
            price = price_value * 100_000
        else:
            price = price_value
        raw_area = clean(source.get("Area_SqFt"))
        area = number(raw_area) if "not published" not in raw_area.casefold() else None
        add(rows, property_id=f"VGN-{clean(source.get('Project_ID'))}", title=title,
            property_type=prop_type, listing_status="Project reference", city=source.get("Location"),
            locality=source.get("Location"), price_inr=price, price_display=raw_price,
            price_basis=source.get("Price_Unit"), area_value=area, area_unit="sq ft" if area else "",
            area_display=f"{raw_area} sq ft" if area else "", bedrooms=source.get("Configuration_BHK"),
            availability=source.get("Available_Flats"), approval_number=source.get("CMDA_Approval"),
            rera_number=source.get("RERA_Number"), amenities=source.get("Amenities"),
            source_name="VGN audited project workbook", source_url=source.get("Source_URL"),
            source_date=source.get("Checked_Date"), date_checked=CHECKED,
            source_status=source.get("Data_Status") or "Project reference; confirm directly",
            notes="Project-level reference; a project price or status does not confirm an available unit. "
                  + clean(source.get("Verification_Notes")))
    book.close()


def export_wisdom(folder: Path, rows: list[dict]):
    book = openpyxl.load_workbook(folder / "Wisdom_Properties_Tamil_Nadu_Master.xlsx", read_only=True, data_only=True)
    sheet = book["01_Project_Master"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for values in iterator:
        source = dict(zip(headers, values))
        title = clean(source.get("Project_Name"))
        prop_type = normalize_type(clean(source.get("Property_Type") or ""))
        if not title or not prop_type:
            continue
        price = number(source.get("Estimated_Min_Plot_Price_INR"))
        raw_size = clean(source.get("Plot_Size_Range_sqft"))
        add(rows, property_id=f"WIS-{clean(source.get('Project_ID'))}", title=title,
            property_type=prop_type, listing_status="Project reference", city=source.get("District"),
            district=source.get("District"), locality=source.get("Location"), price_inr=price,
            price_display=f"₹{price:,.0f} estimated minimum" if price else "", price_basis=source.get("Estimate_Basis"),
            area_display=f"{raw_size} sq ft" if raw_size else "", approval_number=source.get("DTCP_or_Layout_Approval"),
            rera_number=source.get("TN_RERA_Number"), source_name="Wisdom Properties project reference workbook",
            source_date=source.get("Estimate_Source_Date"), date_checked=CHECKED,
            source_url="", source_status=source.get("Estimate_Status") or "Estimate; confirm directly",
            notes="Workbook estimate, not a confirmed offer or available unit. " + clean(source.get("Estimate_Basis")))
    book.close()


def export_cmda_approvals(folder: Path, rows: list[dict]):
    book = openpyxl.load_workbook(folder / "Tamil_Nadu_Property_Market_Master_FINAL.xlsx", read_only=True, data_only=True)
    sheet = book["CMDA_Layout_2024_200"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for values in iterator:
        source = dict(zip(headers, values))
        add(rows, property_id=source.get("Record_ID"), title=f"CMDA layout approval {clean(source.get('Approval_No'))}",
            property_type="Plot", listing_status="Approval record", city="Chennai Metropolitan Area",
            price_display=source.get("Price"), availability=source.get("Availability"),
            approval_number=source.get("Approval_No"), source_name="Tamil Nadu Property Market CMDA approval register",
            source_url=source.get("Source_URL"), source_date="2024", date_checked=CHECKED,
            source_status=source.get("Record_Status") or "Official approval register record",
            notes="Planning approval record only; not a current property-for-sale listing. " + clean(source.get("Notes")))
    book.close()


def import_folder(folder: Path = DEFAULT_SOURCE) -> int:
    if not folder.is_dir():
        raise FileNotFoundError(f"Source folder not found: {folder}")
    rows: list[dict] = []
    export_auction_csv(folder, rows)
    export_nobroker(folder, rows)
    export_stepsstone(folder, rows)
    export_vgn(folder, rows)
    export_wisdom(folder, rows)
    export_cmda_approvals(folder, rows)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import property listings from the local TN Properties folder")
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE)
    args = parser.parse_args()
    count = import_folder(args.source_dir)
    print(f"Imported {count} sanitized property and project records to {OUT}")

